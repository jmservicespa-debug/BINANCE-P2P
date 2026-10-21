import builtins
from pathlib import Path

import pytest
from openpyxl import Workbook

import boletas as b


def test_montos():
    assert b.to_amount(14985) == 14985
    assert b.to_amount(14985.0) == 14985
    assert b.to_amount("14.985") == 14985
    assert b.to_amount("$ 9.990") == 9990
    assert b.to_amount(None) is None
    assert b.to_amount("") is None
    assert b.to_amount(0) is None
    assert b.to_amount(-5) is None
    assert b.to_amount(10.5) is None  # con decimales se revisa a mano
    assert b.to_amount("hola") is None


def libro(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SEGUNDA HOJA"
    for r, (bb, c, d) in {10026: (15000, 15, 14985), 10027: (104530, 105, 104425), 10028: (10000, 10, None), 10029: (1, 1, "x")}.items():
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"] = bb, c, d
    p = tmp_path / "libro.xlsx"
    wb.save(p)
    return p


def test_excel_registro_y_plan(tmp_path):
    rows = b.read_rows(libro(tmp_path), 10026, 10029)
    assert [(r, a) for r, a, _ in rows] == [(10026, 14985), (10027, 104425), (10028, None), (10029, None)]
    log = tmp_path / "log.csv"
    b.record(10026, 14985, "21.938", log)
    todo, notes = b.plan(rows, b.done_rows(log))
    assert todo == [b.Pending(10027, 104425)]  # la ya emitida y las vacías no se emiten
    assert any("ya emitida (folio 21.938)" in n for n in notes)
    assert any("10029" in n and "no válido" in n for n in notes)
    with pytest.raises(SystemExit):
        b.read_rows(tmp_path / "libro.xls", 1, 2)


def test_folio():
    text = "JM SERVICES SPA\nBOLETA EXENTA ELECTRÓNICA NUMERO:\n21.938\nMonto Total\n$ 14.985"
    assert b.FOLIO_RE.search(text).group(1) == "21.938"


# Página de prueba que imita eboleta (la real puede ser distinta: por eso existe el modo diagnóstico).
FAKE = """<!doctype html><meta charset=utf-8><h1>Emitir e-Boleta</h1>
<input type=text id=input-33 readonly><label><input type=checkbox id=input-48> Calculadora</label>
<button type=button>delete</button><button type=button>backspace</button>
<button type=button>7</button><button type=button>8</button><button type=button>9</button><button type=button>4</button>
<button type=button>5</button><button type=button>6</button><button type=button>1</button><button type=button>2</button>
<button type=button>3</button><button type=button>0</button><button type=button id=e>EMITIR</button><div role=dialog id=tipo hidden><button type=button>Afecta</button><button type=button>Exenta</button></div><div id=out></div>
<script>
const box = document.getElementById('input-33');
document.querySelectorAll('button').forEach((b) => b.onclick = () => {
  const t = b.innerText;
  if (/^\\d$/.test(t)) box.value = Number((box.value.replace(/\\D/g, '') || '') + t).toLocaleString('es-CL');
  if (t === 'delete') box.value = '';
  if (t === 'EMITIR') { document.getElementById('tipo').hidden = false; }
  if (t === 'Afecta') document.getElementById('out').innerText = 'BOLETA ELECTRÓNICA AFECTA';
  if (t === 'Exenta') { const m = Number(box.value.replace(/\\D/g, ''));
    document.getElementById('out').innerText = 'BOLETA EXENTA ELECTRÓNICA NUMERO:\\n' + (21938 + m % 7).toLocaleString('es-CL') + '\\nMonto Total $ ' + m.toLocaleString('es-CL'); }
});
</script>"""


def test_emite_en_pagina_de_prueba(tmp_path, monkeypatch):
    from playwright.sync_api import sync_playwright

    page_file = tmp_path / "emitir.html"
    page_file.write_text(FAKE, encoding="utf-8")
    monkeypatch.setattr(b, "EMITIR_URL", page_file.as_uri())
    log = tmp_path / "log.csv"
    monkeypatch.setattr(b, "LOG", log)
    monkeypatch.setattr(builtins, "input", lambda *_: "SI")
    opened = {}

    def fake_open(pw):
        br = pw.chromium.launch(executable_path=_chromium())
        opened["b"] = br
        ctx = br.new_context()
        return ctx, ctx.new_page()

    monkeypatch.setattr(b, "open_browser", fake_open)
    monkeypatch.setattr(b, "ensure_logged_in", lambda page: None)
    b.main(["emitir", "--excel", str(libro(tmp_path)), "--desde", "10026", "--hasta", "10028"])
    rows = b.done_rows(log)
    assert sorted(rows) == [10026, 10027]
    assert rows[10026]["monto"] == "14985" and rows[10026]["folio"]
    # Segunda pasada: nada pendiente, no se emite de nuevo.
    b.main(["emitir", "--excel", str(libro(tmp_path)), "--desde", "10026", "--hasta", "10028"])
    assert len(log.read_text(encoding="utf-8").strip().splitlines()) == 3


def _chromium():
    import glob
    found = sorted(glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome"))
    return found[-1] if found else None


def test_monto_mostrado_como_texto(tmp_path, monkeypatch):
    """Si el SII muestra el monto como texto (no en una casilla), igual se reconoce."""
    from playwright.sync_api import sync_playwright

    html = FAKE.replace("<input type=text id=input-33 readonly>", "<div id=input-33>$ 0</div>").replace("box.value", "box.textContent")
    html = html.replace("if (t === 'delete') box.textContent = '';", "if (t === 'delete') box.textContent = '$ 0';")
    html = html.replace("box.textContent = Number(", "box.textContent = '$ ' + Number(")
    html = html.replace("(box.textContent.replace(/\\\\D/g, '') || '')", "(box.textContent.replace(/\\\\D/g, '').replace(/^0+/, '') || '')")
    f = tmp_path / "e.html"
    f.write_text(html, encoding="utf-8")
    with sync_playwright() as pw:
        br = pw.chromium.launch(executable_path=_chromium())
        page = br.new_page()
        page.goto(f.as_uri())
        b.fill_amount(page, 14985)
        assert "14.985" in page.locator("#input-33").inner_text()
        with pytest.raises(RuntimeError):
            page.evaluate("document.querySelectorAll('button').forEach((x) => { if (/^\\d$/.test(x.innerText)) x.onclick = null; })")
            b.fill_amount(page, 2000)
        br.close()
