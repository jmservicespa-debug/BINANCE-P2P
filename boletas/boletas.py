"""Emite boletas exentas en eboleta.sii.cl con los montos de la columna D del Excel.

Corre en el computador de JM Cambios con Chrome. La clave del SII la escribe la persona en Chrome;
este programa nunca la pide, la ve ni la guarda. Cada boleta emitida se anota en boletas_emitidas.csv
y una fila ya anotada no se vuelve a emitir.

Uso (PowerShell, dentro de la carpeta boletas):
  python boletas.py diagnostico                                      # anota qué tiene la página de emitir
  python boletas.py practica  --excel "C:\\ruta\\libro.xlsx" --desde 10026 --hasta 10030   # llena sin emitir
  python boletas.py emitir    --excel "C:\\ruta\\libro.xlsx" --desde 10026 --hasta 10030
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

HERE = Path(__file__).resolve().parent
EMITIR_URL = "https://eboleta.sii.cl/emitir/"
PROFILE_DIR = HERE / ".perfil-chrome"  # sesión de Chrome solo para este programa (no se sube a GitHub)
LOG = HERE / "boletas_emitidas.csv"
LAST_EXCEL = HERE / "ultimo_excel.txt"  # ruta del último libro usado (solo en este computador)
SHEET = "SEGUNDA HOJA"
COLUMN = "D"
FOLIO_RE = re.compile(r"BOLETA\s+EXENTA\s+ELECTR[OÓ]NICA\s+N[UÚ]MERO\s*:?\s*([\d.]+)", re.I)


@dataclass(frozen=True)
class Pending:
    row: int
    amount: int


# ------------------------------------------------------------------ Excel y registro

def to_amount(value: object) -> int | None:
    """Monto entero en pesos: 14985, 14985.0 o "14.985". Vacío, texto o ≤ 0 → None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            d = Decimal(str(value))
        except InvalidOperation:
            return None
    else:
        t = str(value).strip().replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
        try:
            d = Decimal(t)
        except InvalidOperation:
            return None
    if d <= 0 or d != d.to_integral_value():
        return None  # con decimales no se emite: mejor revisarlo a mano
    return int(d)


def to_number(value: object) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return Decimal(str(value).strip().replace(".", "").replace(",", ".") if isinstance(value, str) else str(value))
    except InvalidOperation:
        return None


def read_rows(path: Path, first: int, last: int, sheet: str = SHEET, column: str = COLUMN) -> list[tuple[int, int | None, object]]:
    """(fila, monto, valor original) de la columna indicada. Lee el valor calculado de las fórmulas (=B-C)."""
    from openpyxl import load_workbook

    if path.suffix.lower() != ".xlsx":
        raise SystemExit("El libro debe ser .xlsx (en Excel: Archivo → Guardar como → Libro de Excel).")
    print(f"Leyendo {path.name}… (un libro grande puede tardar unos segundos)", flush=True)
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f'No encuentro la hoja "{sheet}". Hojas del libro: {", ".join(wb.sheetnames)}')
    ws = wb[sheet]
    from openpyxl.utils import column_index_from_string
    col = column_index_from_string(column)
    # Se leen solo las filas pedidas, de una pasada (leer celda por celda recorre el libro entero cada vez).
    # Columnas B..D: si D viene sin valor guardado, se calcula con su fórmula =B-C.
    out = []
    first_col = min(col, 2)
    cells = ws.iter_rows(min_row=first, max_row=last, min_col=first_col, max_col=col, values_only=True)
    for r, vals in zip(range(first, last + 1), cells):
        raw = vals[col - first_col]
        out.append([r, to_amount(raw), raw, vals[2 - first_col] if col > 2 else None, vals[3 - first_col] if col > 3 else None])
    out += [[r, None, None, None, None] for r in range(first + len(out), last + 1)]  # filas más allá del final del libro
    wb.close()
    missing = [x for x in out if x[2] in (None, "")]
    if missing:
        # Fórmula sin resultado guardado (pasa si el libro no lo calculó al guardar): se resuelve =B{n}-C{n}.
        wf = load_workbook(path, data_only=False, read_only=True)
        rows_f = wf[sheet].iter_rows(min_row=first, max_row=last, min_col=col, max_col=col, values_only=True)
        formulas = {r: (f or "") for r, (f,) in zip(range(first, last + 1), rows_f)}
        wf.close()
        for x in missing:
            m = re.fullmatch(r"=\s*B(\d+)\s*-\s*C(\d+)\s*", str(formulas.get(x[0], "")), re.I)
            if m and int(m.group(1)) == int(m.group(2)) == x[0]:
                bv, cv = to_number(x[3]), to_number(x[4])
                if bv is not None and cv is not None:
                    x[2] = f"=B{x[0]}-C{x[0]} → {bv - cv}"
                    x[1] = to_amount(bv - cv)
    return [(r, amount, raw) for r, amount, raw, _, _ in out]


def pick_excel(given: str | None) -> Path:
    """Ruta del libro: la que se escribió, o se elige en una ventana (y se recuerda para la próxima vez)."""
    if given and given.strip().upper() != "RUTA":
        path = Path(given.strip().strip('"'))
    else:
        last = LAST_EXCEL.read_text(encoding="utf-8").strip() if LAST_EXCEL.exists() else ""
        if last and Path(last).exists() and input(f'Usar el libro "{last}"? (Enter = sí, N = elegir otro): ').strip().upper() != "N":
            path = Path(last)
        else:
            try:
                import tkinter as tk
                from tkinter import filedialog
                root = tk.Tk()
                root.withdraw()
                root.attributes("-topmost", True)
                chosen = filedialog.askopenfilename(title="Elige tu libro de Excel", filetypes=[("Excel", "*.xlsx"), ("Todos", "*.*")])
                root.destroy()
            except Exception:
                chosen = input("Pega la ruta del Excel (clic derecho en el archivo → Copiar como ruta de acceso): ")
            if not chosen:
                raise SystemExit("No elegiste ningún archivo.")
            path = Path(chosen.strip().strip('"'))
    if not path.exists():
        raise SystemExit(f"No encuentro el archivo: {path}")
    LAST_EXCEL.write_text(str(path), encoding="utf-8")
    return path


def done_rows(log: Path | None = None) -> dict[int, dict]:
    log = log or LOG
    if not log.exists():
        return {}
    with log.open(newline="", encoding="utf-8") as f:
        return {int(x["fila"]): x for x in csv.DictReader(f)}


def record(row: int, amount: int, folio: str, log: Path | None = None) -> None:
    log = log or LOG
    new = not log.exists()
    with log.open("a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(["fila", "monto", "folio", "fecha_hora"])
        w.writerow([row, amount, folio, datetime.now().isoformat(timespec="seconds")])


def plan(rows, done) -> tuple[list[Pending], list[str]]:
    """Separa lo que se emite de lo que se salta (ya emitida, vacía o con monto raro)."""
    todo, notes = [], []
    for r, amount, raw in rows:
        if r in done:
            notes.append(f"fila {r}: ya emitida (folio {done[r]['folio']})")
        elif amount is None:
            if raw not in (None, ""):
                notes.append(f"fila {r}: monto no válido ({raw!r}), se salta")
        else:
            todo.append(Pending(r, amount))
    return todo, notes


def fmt(n: int) -> str:
    return f"{n:,}".replace(",", ".")


# ------------------------------------------------------------------ Chrome

def open_browser(pw):
    PROFILE_DIR.mkdir(exist_ok=True)
    try:
        ctx = pw.chromium.launch_persistent_context(PROFILE_DIR, channel="chrome", headless=False, viewport=None)
    except Exception:
        ctx = pw.chromium.launch_persistent_context(PROFILE_DIR, headless=False, viewport=None)
    page = ctx.pages[0] if ctx.pages else ctx.new_page()
    return ctx, page


def ensure_logged_in(page) -> None:
    page.goto(EMITIR_URL)
    input("\nSi Chrome pide entrar al SII, entra TÚ con tu RUT y clave. Cuando veas la pantalla de emitir boleta, presiona Enter aquí… ")
    if "eboleta.sii.cl/emitir" not in page.url:
        page.goto(EMITIR_URL)


def first_visible(locators, timeout: int = 1500):
    for loc in locators:
        try:
            loc.first.wait_for(state="visible", timeout=timeout)
            return loc.first
        except Exception:
            continue
    return None


def choose_exenta(page) -> bool:
    """Marca "Exenta" si la pantalla tiene esa opción. En eboleta (oct. 2026) no aparece en la pantalla de emitir:
    el tipo lo confirma la boleta emitida, que debe decir BOLETA EXENTA (si no, el programa se detiene)."""
    re_ex = re.compile(r"exent", re.I)
    for sel in page.locator("select").all():
        options = sel.locator("option").all_inner_texts()
        match = next((o for o in options if re_ex.search(o)), None)
        if match:
            sel.select_option(label=match)
            return True
    el = first_visible([page.get_by_role("radio", name=re_ex), page.get_by_label(re_ex), page.get_by_role("tab", name=re_ex),
                        page.get_by_role("button", name=re_ex)], timeout=800)
    if el is None:
        return False
    if el.get_attribute("type") in ("radio", "checkbox"):
        el.check()
    else:
        el.click()
    return True


def amount_on_screen(page, amount: int, before: str) -> bool:
    """¿La pantalla muestra el monto? Sirve una casilla con ese número, o que "14.985" aparezca en el texto
    de la página después de marcarlo (y no estuviera antes)."""
    for box in page.locator("input[type=text], input[type=tel], input[type=number]").all():
        try:
            if re.sub(r"\D", "", box.input_value()) == str(amount):
                return True
        except Exception:
            continue
    pat = re.compile(r"(?<![\d.])" + re.escape(fmt(amount)) + r"(?![\d.])") if amount >= 1000 else re.compile(r"\$\s*" + str(amount) + r"(?!\d)")
    after = page.locator("body").inner_text()
    return len(pat.findall(after)) > len(pat.findall(before))


def fill_amount(page, amount: int) -> None:
    """eboleta usa un teclado en pantalla (0–9, delete, backspace): se borra y se marcan los dígitos.
    Después se comprueba que la casilla muestre exactamente el monto; si no, no se emite."""
    key = lambda name: page.get_by_role("button", name=name, exact=True)
    before = page.locator("body").inner_text()
    if key("0").count():
        if key("delete").count():
            key("delete").first.click()
        for ch in str(amount):
            key(ch).first.click()
    else:
        box = page.locator("input[type=text]").first
        if not box.count():
            raise RuntimeError("no encontré dónde escribir el monto")
        box.fill(str(amount))
    page.wait_for_timeout(300)
    if not amount_on_screen(page, amount, before):
        raise RuntimeError(f"no veo el monto {fmt(amount)} en la pantalla después de marcarlo")


def pick_option(page, opener: re.Pattern, choice: re.Pattern, what: str) -> None:
    """Selector desplegable de eboleta (Vuetify): se abre el que muestra `opener` y se elige la opción `choice`."""
    box = first_visible([page.get_by_role("button", name=opener), page.locator("[role=button]").filter(has_text=opener)], timeout=5000)
    if box is None:
        raise RuntimeError(f"no encontré el selector de {what}")
    if choice.search(box.inner_text()):
        return  # ya estaba elegido
    box.click()
    opt = first_visible([page.get_by_role("option", name=choice), page.locator(".v-list-item, [role=option], [role=listitem]").filter(has_text=choice)], timeout=4000)
    if opt is None:
        raise RuntimeError(f"no encontré la opción de {what} que corresponde")
    opt.click()
    page.wait_for_timeout(300)
    if not choice.search(box.inner_text()):
        raise RuntimeError(f"el selector de {what} no quedó en la opción correcta")


def pick_exenta_after_emit(page) -> None:
    """Tras apretar EMITIR, eboleta abre una ventana con el tipo de boleta ("Boleta afecta ▾"): se cambia a exenta."""
    pick_option(page, re.compile(r"boleta\s+(afecta|exenta)", re.I), re.compile(r"boleta\s+exenta", re.I), "tipo de boleta")


def open_dialog(page) -> None:
    """Aprieta EMITIR de la pantalla y deja la ventana del tipo en "Boleta exenta" (sin confirmar)."""
    btn = first_visible([page.get_by_role("button", name=re.compile(r"^\s*emitir\s*$", re.I)), page.get_by_text(re.compile(r"^\s*emitir\s*$", re.I))])
    if btn is None:
        raise RuntimeError('no encontré el botón "Emitir"')
    btn.click()
    pick_exenta_after_emit(page)


def dialog_buttons(page) -> list[str]:
    scope = first_visible([page.locator(".v-dialog--active"), page.locator("[role=dialog]"), page.locator(".v-dialog"), page.locator(".v-overlay__content")], timeout=1000) or page
    return [t.strip() for t in scope.get_by_role("button").all_inner_texts() if t.strip()]


def emit(page, before_final=lambda: None) -> None:
    """EMITIR de la pantalla → ventana en "Boleta exenta" → EMITIR de la ventana (el método de pago se deja vacío).
    `before_final` se llama justo antes del clic final: ahí la fila se anota, para que nunca se emita dos veces."""
    open_dialog(page)
    re_final = re.compile(r"^\s*emitir\s*$", re.I)
    dialog = first_visible([page.locator(".v-dialog--active"), page.locator("[role=dialog]"), page.locator(".v-dialog"), page.locator(".v-overlay__content")], timeout=1000)
    if dialog is not None:
        final = first_visible([dialog.get_by_role("button", name=re_final)], timeout=3000)
    else:
        # Sin ventana reconocible: el EMITIR de la ventana es el último de la página (el primero es el de atrás).
        buttons = page.get_by_role("button", name=re_final)
        final = buttons.last if buttons.count() >= 2 else None
    if final is None:
        raise RuntimeError('no encontré el botón EMITIR de la ventana')
    before_final()
    final.click()


def folio_in_json(obj) -> str | None:
    """Busca un campo "folio" (o parecido) en la respuesta del SII."""
    if isinstance(obj, dict):
        for k, v in obj.items():
            if re.fullmatch(r"(folio|nro_?folio|numero_?folio|folioboleta)", str(k), re.I) and str(v).strip().replace(".", "").isdigit():
                return str(v).strip()
        obj = list(obj.values())
    if isinstance(obj, list):
        for v in obj:
            got = folio_in_json(v)
            if got:
                return got
    return None


def read_folio(page, amount: int, responses: list, timeout_s: int = 30) -> str | None:
    """Número de la boleta emitida. El SII la muestra como imagen, así que se busca en el texto de la página y sus
    marcos, y en lo que respondió el SII al emitir. Si se ve un folio con otro monto, se detiene."""
    deadline = datetime.now().timestamp() + timeout_s
    while datetime.now().timestamp() < deadline:
        for frame in page.frames:
            try:
                text = frame.locator("body").inner_text(timeout=1000)
            except Exception:
                continue
            m = FOLIO_RE.search(text)
            if m:
                if fmt(amount) not in text and str(amount) not in text:
                    raise RuntimeError(f"la boleta {m.group(1)} no muestra el monto {fmt(amount)}: revísala en el SII")
                return m.group(1)
        for r in list(responses):
            try:
                got = folio_in_json(r.json())
            except Exception:
                got = None
            if got:
                return got
        page.wait_for_timeout(500)
    return None


# ------------------------------------------------------------------ comandos

DUMP_JS = """() => {
  const vis = (el) => !!(el.offsetWidth || el.offsetHeight || el.getClientRects().length);
  const lab = (el) => (el.labels && el.labels[0] ? el.labels[0].innerText : '') || el.getAttribute('aria-label') || el.placeholder || '';
  const out = [];
  document.querySelectorAll('input, select, textarea, button, [role=button], [role=radio], [role=tab], [role=option], [role=listitem], .v-list-item, a').forEach((el) => {
    if (!vis(el) || el.type === 'password' || el.type === 'hidden') return;
    const val = 'value' in el && el.tagName !== 'BUTTON' && el.value ? ` valor="${String(el.value).slice(0, 30)}"` : '';
    const opts = el.tagName === 'SELECT' ? ' opciones=' + [...el.options].map((o) => o.text).join(' | ') : '';
    out.push(`${el.tagName.toLowerCase()} role=${el.getAttribute('role') || ''} type=${el.type || ''} id=${el.id} texto="${(el.innerText || '').trim().replace(/\\s+/g, ' ').slice(0, 50)}" etiqueta="${lab(el).trim().slice(0, 40)}"${val}${opts}`);
  });
  document.querySelectorAll('[role=dialog], .v-dialog, .v-overlay__content, .modal').forEach((d) => {
    if (vis(d)) out.push('VENTANA: ' + d.innerText.trim().replace(/\\s+/g, ' ').slice(0, 300));
  });
  return out;
}"""


def cmd_diagnostico(args) -> None:
    from playwright.sync_api import sync_playwright

    with sync_playwright() as pw:
        ctx, page = open_browser(pw)
        ensure_logged_in(page)
        problem = None
        if args.tras_emitir:
            # Marca $1.000 y aprieta EMITIR solo para ver la ventana del tipo de boleta; NO elige nada.
            try:
                fill_amount(page, 1000)
                btn = first_visible([page.get_by_role("button", name=re.compile(r"^\s*emitir\s*$", re.I))])
                if btn is None:
                    raise RuntimeError('no encontré el botón "EMITIR"')
                btn.click()
                page.wait_for_timeout(2500)
            except Exception as e:  # se anota igual lo que hay en pantalla, sin apretar EMITIR
                problem = str(e)
        info = page.evaluate(DUMP_JS)
        if problem:
            info.insert(0, f"PROBLEMA: {problem} (no se apretó EMITIR)")
        page.screenshot(path=str(HERE / "diagnostico_eboleta.png"))
        target = HERE / "diagnostico_eboleta.txt"
        target.write_text(f"{page.url}\n" + "\n".join(info), encoding="utf-8")
        print(f"\nListo: {len(info)} elementos anotados en {target} (y foto de la pantalla en diagnostico_eboleta.png)")
        if problem:
            print(f"Ojo: {problem}. No se apretó EMITIR.")
        print("Mándale a Claude una foto o el texto de ese archivo (no contiene tu clave).")
        msg = "NO elijas nada en la ventana de Chrome; sácale foto. " if args.tras_emitir else ""
        input(f"\n{msg}Chrome queda abierto: presiona Enter aquí para cerrarlo… ")
        ctx.close()


def cmd_run(args, really_emit: bool) -> None:
    from playwright.sync_api import sync_playwright

    rows = read_rows(pick_excel(args.excel), args.desde, args.hasta, args.hoja, args.columna)
    todo, notes = plan(rows, done_rows())
    for n in notes:
        print("  ·", n)
    if not todo:
        empty = [r for r, amount, raw in rows if raw in (None, "")]
        print("No hay boletas pendientes en ese rango.")
        if empty:
            print(f"  La columna {args.columna} está vacía en: {', '.join(map(str, empty[:20]))}{'…' if len(empty) > 20 else ''}"
                  f" (hoja \"{args.hoja}\"). Revisa el número de fila en Excel.")
        return
    total = sum(p.amount for p in todo)
    print(f"\n{len(todo)} boletas exentas, total ${fmt(total)}:")
    for p in todo:
        print(f"  fila {p.row}: ${fmt(p.amount)}")
    if really_emit:
        if input(f"\n¿Emitir estas {len(todo)} boletas en el SII? Escribe SI para continuar: ").strip().upper() != "SI":
            print("Cancelado. No se emitió nada.")
            return
    else:
        print("\nMODO PRÁCTICA: marca el monto, abre la ventana y elige Boleta exenta, pero NO aprieta el botón final.")

    with sync_playwright() as pw:
        ctx, page = open_browser(pw)
        ensure_logged_in(page)
        emitted = 0
        for p in todo:
            try:
                page.goto(EMITIR_URL)
                page.wait_for_load_state("domcontentloaded")
                exenta = choose_exenta(page)
                fill_amount(page, p.amount)
                if not really_emit:
                    open_dialog(page)  # abre la ventana y deja "Boleta exenta", pero NO aprieta el botón final
                    print(f"fila {p.row}: ${fmt(p.amount)} marcado y ventana en Boleta exenta. Botones de la ventana: {', '.join(dialog_buttons(page))}")
                    input("  Revisa Chrome (NO apretes nada en la ventana) y presiona Enter aquí para seguir… ")
                    close = first_visible([page.get_by_role("button", name=re.compile(r"^\s*(close|cerrar|cancelar)\s*$", re.I))], timeout=1000)
                    if close is not None:
                        close.click()
                    continue
                responses = []

                def keep(r, responses=responses):
                    if "sii.cl" in r.url and "json" in (r.headers.get("content-type") or ""):
                        responses.append(r)

                page.on("response", keep)
                try:
                    emit(page, before_final=lambda: record(p.row, p.amount, "pendiente"))
                    folio = read_folio(page, p.amount, responses)
                finally:
                    page.remove_listener("response", keep)
                emitted += 1
                if folio:
                    record(p.row, p.amount, folio)
                    print(f"  ✓ fila {p.row}: ${fmt(p.amount)} → boleta N° {folio}")
                    continue
                # Emitida pero sin número legible: queda anotada (no se repite) y se guarda foto para revisar.
                shot = HERE / f"boleta_fila_{p.row}.png"
                page.screenshot(path=str(shot))
                record(p.row, p.amount, "emitida-sin-numero")
                info = page.evaluate(DUMP_JS)
                (HERE / "diagnostico_boleta_emitida.txt").write_text(
                    "\n".join([page.url, f"marcos: {[f.url for f in page.frames]}", f"respuestas: {[r.url for r in responses]}"] + info), encoding="utf-8")
                print(f"  ✓ fila {p.row}: ${fmt(p.amount)} emitida, pero no pude leer el número (foto en {shot.name}).")
                print("    Queda anotada y no se repetirá. Me detengo para que lo revises; mándale a Claude diagnostico_boleta_emitida.txt.")
                break
            except Exception as e:  # cualquier cosa rara: se detiene para no emitir de más
                print(f"\n✗ fila {p.row}: {e}. Me detengo aquí; revisa el SII antes de seguir.")
                input("Presiona Enter para cerrar Chrome… ")
                break
        if really_emit:
            print(f"\nEmitidas: {emitted} de {len(todo)}. Registro: {LOG}")
        ctx.close()


def cmd_marcar(args) -> None:
    """Anota a mano una fila ya emitida (por ejemplo, hecha en el SII sin el programa) para que no se repita."""
    record(args.fila, args.monto, args.folio)
    print(f"Anotada: fila {args.fila}, ${fmt(args.monto)}, boleta N° {args.folio}. No se volverá a emitir.")


def cmd_revisar(args) -> None:
    """Muestra lo que hay en el archivo guardado (columnas A–E, valor y fórmula). No abre Chrome ni emite nada."""
    from openpyxl import load_workbook

    path = pick_excel(args.excel)
    print(f"Leyendo {path.name}…", flush=True)
    wv = load_workbook(path, data_only=True, read_only=True)
    wf = load_workbook(path, data_only=False, read_only=True)
    print(f"Archivo: {path}  ·  guardado: {datetime.fromtimestamp(path.stat().st_mtime):%d-%m-%Y %H:%M}")
    print("Hojas: " + " | ".join(f"{n} ({wv[n].max_row} filas)" for n in wv.sheetnames))
    if args.hoja not in wv.sheetnames:
        print(f'No existe la hoja "{args.hoja}".')
        return
    vals = wv[args.hoja].iter_rows(min_row=args.desde, max_row=args.hasta, max_col=5, values_only=True)
    forms = wf[args.hoja].iter_rows(min_row=args.desde, max_row=args.hasta, max_col=5, values_only=True)
    print("fila  | A | B | C | D (valor) | D (fórmula) | E")
    for r, v, f in zip(range(args.desde, args.hasta + 1), vals, forms):
        v = list(v) + [None] * (5 - len(v))
        f = list(f) + [None] * (5 - len(f))
        print(f"{r} | {v[0]} | {v[1]} | {v[2]} | {v[3]} | {f[3] if isinstance(f[3], str) and f[3].startswith('=') else ''} | {v[4]}")
    wv.close()
    wf.close()


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description="Boletas exentas en eboleta.sii.cl desde el Excel (columna D).")
    sub = ap.add_subparsers(dest="cmd", required=True)
    d = sub.add_parser("diagnostico", help="anota los campos de la página de emitir")
    d.add_argument("--tras-emitir", action="store_true", help="marca $1.000, aprieta EMITIR y anota la ventana del tipo de boleta (sin elegir)")
    mk = sub.add_parser("marcar", help="anota a mano una fila ya emitida para que no se repita")
    mk.add_argument("--fila", type=int, required=True)
    mk.add_argument("--monto", type=int, required=True)
    mk.add_argument("--folio", required=True)
    rv = sub.add_parser("revisar", help="muestra qué hay en esas filas del archivo guardado (no emite)")
    rv.add_argument("--excel")
    rv.add_argument("--desde", type=int, required=True)
    rv.add_argument("--hasta", type=int, required=True)
    rv.add_argument("--hoja", default=SHEET)
    for name in ("practica", "emitir"):
        p = sub.add_parser(name)
        p.add_argument("--excel", help="ruta del libro .xlsx (si no se pone, se elige en una ventana)")
        p.add_argument("--desde", type=int, required=True, help="primera fila del Excel")
        p.add_argument("--hasta", type=int, required=True, help="última fila del Excel")
        p.add_argument("--hoja", default=SHEET)
        p.add_argument("--columna", default=COLUMN)
    args = ap.parse_args(argv)
    if args.cmd == "diagnostico":
        cmd_diagnostico(args)
    elif args.cmd == "marcar":
        cmd_marcar(args)
    elif args.cmd == "revisar":
        cmd_revisar(args)
    else:
        if args.hasta < args.desde or args.hasta - args.desde > 500:
            sys.exit("Rango de filas no válido (máximo 500 por tanda).")
        cmd_run(args, really_emit=args.cmd == "emitir")


if __name__ == "__main__":
    main()
